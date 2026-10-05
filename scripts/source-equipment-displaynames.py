"""Read-only ItemInfo label evidence; never edits the input workbooks."""
import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

pattern = re.compile(
    r"GearFootwear_|NeckWear_Turtleneck_Tight[2-7]$|GearNeckpad_|FaceGear_|G_CompressionT_|Gear_Undershirt_|"
    r"GearMouthpiece_|"
    r"ElbowGear_RubberBands[1-3]$|Japanese_08|"
    r"ArmTattoo_(?:24|25|34|35)$|CujoMatty_ArmTats_(?:24|25|34|35)$|LegTattoo_TEST$|"
    r"GearArmSleeve_NikeProDriFitSleeve2|"
    r"ThighPad_(?:Shells|BeastMode_Logo|CheatCode_Logo|Chosen_Logo|Joker_Logo|Lion_Logo)$|Towel2_South$",
    re.IGNORECASE,
)
results = []
for source in sys.argv[1:]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook:
        if not re.search(r"iteminfo|equipment", sheet.title, re.IGNORECASE):
            continue
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        fields = {str(value): index for index, value in enumerate(header) if value is not None}
        if "ItemName" not in fields:
            continue
        matches = []
        for row_number, row in enumerate(rows, 2):
            name = row[fields["ItemName"]]
            image = row[fields["ImageId"]] if "ImageId" in fields else None
            image_match = isinstance(image, str) and re.search(r"CM_ArmTattoo_(?:24|25|34|35)$", image)
            if not isinstance(name, str) or not (pattern.search(name) or image_match):
                continue
            matches.append({"row": row_number, **{
                field: row[fields[field]] for field in
                ["ItemName", "DisplayName", "ItemId", "CharacterRole", "LoadoutSlots", "ItemTags", "ImageId"]
                if field in fields
            }})
        sheets.append({"sheet": sheet.title, "headers": fields, "matches": matches})
    results.append({"source": str(Path(source).resolve()), "sheets": sheets})
    workbook.close()
print(json.dumps(results, indent=2, ensure_ascii=False))
